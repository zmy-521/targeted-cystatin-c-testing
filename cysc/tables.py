"""Baseline tables and a workbook collecting the manuscript's aggregate sources."""
import numpy as np
import pandas as pd
from scipy.stats import mannwhitneyu,chi2_contingency,fisher_exact
from .io import FROZEN,DICTIONARY,output,load_private,save_table,require,checkpoint
from .preprocess import REFERENCE

def baseline(d):
    units={r['name']:r['unit'] for r in DICTIONARY}
    units.update({'sex':'—','CysC_mg_L':'mg/L','eGFRcys_2012':'mL/min/1.73 m²','eGFR_difference':'mL/min/1.73 m²','eGFR_ratio':'ratio'})
    # Same approved clinical order for development and primary external.
    features=FROZEN['full'][:2]+FROZEN['full'][3:]+['eGFRcr_2021']+REFERENCE
    rows=[]
    for f in features:
        x=d.loc[d.D2==0,f].dropna();y=d.loc[d.D2==1,f].dropna()
        if f=='sex':
            counts=np.array([[int(x.sum()),len(x)-int(x.sum())],[int(y.sum()),len(y)-int(y.sum())]])
            if np.any(counts.sum(axis=0)==0):p=1.;test='Fisher exact'
            else:
                _,p,_,expected=chi2_contingency(counts,correction=False)
                test='Chi-square'
                if (expected<5).any():_,p=fisher_exact(counts);test='Fisher exact'
            fmt=lambda a:f'{int(a.sum())} ({100*a.mean():.1f}%)'
        else:
            p=mannwhitneyu(x,y,alternative='two-sided',method='asymptotic',use_continuity=True).pvalue if len(x) and len(y) else np.nan
            test='Mann–Whitney U';fmt=lambda a:f'{a.median():.2f} [{a.quantile(.25):.2f}, {a.quantile(.75):.2f}]' if len(a) else 'NA'
        part='Phenotype characterization / reference measurements' if f in REFERENCE else 'Renal core' if f=='eGFRcr_2021' else 'Demographics' if f in ['age','sex'] else 'Hematology' if f in FROZEN['full'][3:14] else 'Biochemistry'
        rows.append({'Part':part,'Variable':f,'Unit':units[f],'D2_negative':fmt(x),'D2_positive':fmt(y),'Negative_available_n':len(x),'Positive_available_n':len(y),'P_value':p,'Test':test})
    return pd.DataFrame(rows)

def three_line_docx(path,t,title,n0,n1):
    from docx import Document
    from docx.shared import Pt,Inches
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10)
    doc.add_paragraph(title).runs[0].bold=True
    tb=doc.add_table(rows=1,cols=5)
    for c,s in zip(tb.rows[0].cells,['Characteristic','Unit',f'D2-negative (n = {n0})',f'D2-positive (n = {n1})','P value']):c.text=s
    last=None
    for _,r in t.iterrows():
        if r.Part!=last:
            cells=tb.add_row().cells;cells[0].merge(cells[-1]);cells[0].text=r.Part;cells[0].paragraphs[0].runs[0].bold=True;last=r.Part
        cells=tb.add_row().cells
        p='NA' if pd.isna(r.P_value) else '<0.001' if r.P_value<.001 else f'{r.P_value:.3f}'
        for c,v in zip(cells,[r.Variable,r.Unit,r.D2_negative,r.D2_positive,p]):c.text=str(v)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'8');borders.append(e)
    tb._tbl.tblPr.append(borders)
    header=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(header)
    for cell in tb.rows[0].cells:
        b=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');b.append(e);cell._tc.get_or_add_tcPr().append(b)
    for row in tb.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for p in c.paragraphs:p.paragraph_format.space_after=Pt(2)
    doc.add_paragraph('Values are median [Q1, Q3] or n (%), using available observations without imputation. P values are two-sided. Sex was coded as 0 = female and 1 = male.')
    doc.add_paragraph('Cystatin C and cystatin C-derived measurements were used for phenotype definition/characterization and were not included in the machine-learning predictor set. SCr was excluded because of structural redundancy with eGFRcr.')
    doc.save(path)

def run(cfg):
    require(cfg,'profiles');p=output(cfg)/'tables'
    for co,name,title in [('development','Main_Table_1','Baseline characteristics of the development cohort'),('external','Table_S2','Baseline characteristics of the external validation cohort')]:
        d=load_private(cfg,co);t=baseline(d);save_table(cfg,name,t)
        three_line_docx(p/(name+'.docx'),t,title,int((d.D2==0).sum()),int(d.D2.sum()))
    names={'Table1':'Main_Table_1','S1':'Table_S1_variable_eligibility','S2':'Table_S2','S3_models':'Full_model_comparison',
           'S4A_reduction':'Feature_reduction','S4B_stability':'Panel_stability_repeats','S4B_summary':'Panel_stability_summary',
           'S5_performance':'Performance_CI','S5_paired_differences':'Paired_differences','S6_calibration':'Calibration','S6_crossfit_parameters':'Crossfit_parameters',
           'S7_testing':'Targeted_testing','S8_binary_phenotypes':'Phenotype_binary_comparisons','S9_continuous_phenotypes':'Phenotype_continuous_comparisons'}
    with pd.ExcelWriter(p/'Manuscript_aggregate_sources.xlsx',engine='openpyxl') as w:
        for sheet,file in names.items():pd.read_csv(p/(file+'.csv')).to_excel(w,sheet_name=sheet,index=False)
    checkpoint(cfg,'tables')
